"""Compare older ZSE records (bevennyamande) against this repo's local archive.

Implements the five handoff tasks in DATA_SOURCES.md:
  1. Import older ZSE records (2020-10-08 through 2024-11-08)
  2. Normalize into archive-single-file/ shape
  3. Compare overlap from 2024-11-09 onward
  4. Flag inconsistencies by date, company, field, local value, older value
  5. Emit reviewable reports before any archive changes
"""

from __future__ import annotations

import csv
import json
import os
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parent.parent
OLDER_DIR = REPO_ROOT / "reference-data" / "older-zse"
OLDER_CSV_DIR = OLDER_DIR / "csv-daily-pricesheets"
OLDER_XLSX_DIR = OLDER_DIR / "xls-daily-price-sheets"
LOCAL_ARCHIVE_DIR = REPO_ROOT / "archive"
LOCAL_SINGLE_FILE_DIR = REPO_ROOT / "archive-single-file"
NORMALIZED_DIR = REPO_ROOT / "comparison" / "older-normalized"
REPORTS_DIR = REPO_ROOT / "comparison" / "reports"

BRIEF_IMPORT_START = date(2020, 10, 8)
BRIEF_IMPORT_END = date(2024, 11, 8)
BRIEF_LOCAL_START = date(2024, 11, 9)
NUMERIC_TOLERANCE = 1e-4

FIELDS = ("open", "close", "volume")
FIELD_TO_INDEX = {"open": 0, "close": 1, "volume": 2}
FIELD_DISPLAY = {"open": "opening_price", "close": "closing_price", "volume": "volume"}


@dataclass(frozen=True)
class DailyRow:
    company: str
    open: float | None
    close: float | None
    volume: float | None

    def with_field(self, name: str, value: float | None) -> "DailyRow":
        kwargs = {"open": self.open, "close": self.close, "volume": self.volume}
        kwargs[name] = value
        return DailyRow(self.company, **kwargs)


@dataclass
class DaySnapshot:
    rows: dict[str, DailyRow] = field(default_factory=dict)
    duplicates: list[dict[str, DailyRow]] = field(default_factory=list)


@dataclass
class Issue:
    date: str
    company: str
    field: str
    local_value: str
    older_value: str
    issue_type: str
    note: str = ""


def parse_csv_filename(name: str) -> date | None:
    stem = Path(name).stem
    try:
        return datetime.strptime(stem, "%m-%d-%Y").date()
    except ValueError:
        return None


def parse_xlsx_filename(name: str) -> date | None:
    return parse_csv_filename(name)


def _is_company_column(value: str) -> bool:
    s = value.strip()
    if not s:
        return False
    if s.lower() in ("equities", "0"):
        return False
    if s.replace(".", "").replace("-", "").isdigit():
        return False
    return True


def _parse_csv_data_row(row: list[str]) -> tuple[str, float | None, float | None, float | None] | None:
    if len(row) < 5:
        return None
    if _is_company_column(row[1]):
        company_idx, op_idx, cl_idx, vol_idx = 1, 2, 3, 4
    elif len(row) >= 6 and _is_company_column(row[2]):
        company_idx, op_idx, cl_idx, vol_idx = 2, 3, 4, 5
    else:
        return None
    company = row[company_idx].strip()
    try:
        op = float(row[op_idx]) if row[op_idx].strip() else None
        cl = float(row[cl_idx]) if row[cl_idx].strip() else None
        vol = float(row[vol_idx]) if row[vol_idx].strip() else None
    except (ValueError, IndexError):
        return None
    return company, op, cl, vol


def load_older_csv(path: Path) -> DaySnapshot:
    snap = DaySnapshot()
    with path.open(newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            parsed = _parse_csv_data_row(row)
            if parsed is None:
                continue
            company, op, cl, vol = parsed
            if company in snap.rows:
                snap.duplicates.append({company: snap.rows[company]})
            snap.rows[company] = DailyRow(company, op, cl, vol)
    return snap


def load_older_xlsx(path: Path) -> DaySnapshot:
    snap = DaySnapshot()
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError):
        return snap
    try:
        ws = wb.active
        if ws is None:
            return snap
        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 4:
                continue
            company = row[1]
            if not isinstance(company, str) or not company.strip():
                continue
            if company.strip().lower() == "equities":
                continue
            try:
                op = float(row[2]) if row[2] is not None and str(row[2]).strip() else None
                cl = float(row[3]) if row[3] is not None and str(row[3]).strip() else None
            except (ValueError, TypeError):
                continue
            vol = None
            if len(row) > 4 and row[4] is not None and str(row[4]).strip():
                try:
                    vol = float(row[4])
                except (ValueError, TypeError):
                    vol = None
            if company in snap.rows:
                snap.duplicates.append({company: snap.rows[company]})
            snap.rows[company] = DailyRow(company.strip(), op, cl, vol)
    finally:
        wb.close()
    return snap


def load_older_data() -> dict[str, DaySnapshot]:
    out: dict[str, DaySnapshot] = {}

    for path in sorted(OLDER_CSV_DIR.glob("*.csv")):
        d = parse_csv_filename(path.name)
        if d is None:
            continue
        key = d.isoformat()
        snap = load_older_csv(path)
        if snap.rows:
            out[key] = merge_snapshots(out.get(key), snap)

    for path in sorted(OLDER_XLSX_DIR.glob("*.xlsx")):
        d = parse_xlsx_filename(path.name)
        if d is None:
            continue
        key = d.isoformat()
        snap = load_older_xlsx(path)
        if snap.rows:
            out[key] = merge_snapshots(out.get(key), snap)

    return out


def merge_snapshots(existing: DaySnapshot | None, incoming: DaySnapshot) -> DaySnapshot:
    if existing is None:
        return incoming
    for company, row in incoming.rows.items():
        if company in existing.rows:
            existing.duplicates.append({company: existing.rows[company]})
        existing.rows[company] = row
    existing.duplicates.extend(incoming.duplicates)
    return existing


def _is_real_company(name: str) -> bool:
    return bool(name) and name == name.strip() and name.upper() != name


def _parse_archive_value(raw: Any) -> float | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def load_local_from_consolidated() -> dict[str, DaySnapshot]:
    out: dict[str, DaySnapshot] = {}
    files = {
        "open": LOCAL_SINGLE_FILE_DIR / "open_price_zse.json",
        "close": LOCAL_SINGLE_FILE_DIR / "close_price_zse.json",
        "vol": LOCAL_SINGLE_FILE_DIR / "vol_traded_zse.json",
    }
    data_by_field: dict[str, dict[str, dict[str, float | None]]] = {}
    for label, path in files.items():
        if not path.exists():
            continue
        with path.open() as f:
            doc = json.load(f)
        cols = doc.get("columns", [])
        idx = doc.get("index", [])
        data = doc.get("data", [])
        per_day: dict[str, dict[str, float | None]] = {}
        for di, dlabel in enumerate(idx):
            day = str(dlabel).split("T")[0]
            per_day[day] = {}
            for ci, company in enumerate(cols):
                if not _is_real_company(company):
                    continue
                if di >= len(data) or ci >= len(data[di]):
                    continue
                per_day[day][company] = _parse_archive_value(data[di][ci])
        data_by_field[label] = per_day

    all_days = set()
    for d in data_by_field.values():
        all_days.update(d.keys())

    for day in all_days:
        snap = DaySnapshot()
        opens = data_by_field.get("open", {}).get(day, {})
        closes = data_by_field.get("close", {}).get(day, {})
        vols = data_by_field.get("vol", {}).get(day, {})
        for company in set(opens) | set(closes) | set(vols):
            snap.rows[company] = DailyRow(
                company,
                opens.get(company),
                closes.get(company),
                vols.get(company),
            )
        out[day] = snap
    return out


def load_local_from_folders() -> dict[str, list[DaySnapshot]]:
    out: dict[str, list[DaySnapshot]] = defaultdict(list)
    for folder in sorted(LOCAL_ARCHIVE_DIR.iterdir()):
        if not folder.is_dir():
            continue
        day = folder.name.split(" ")[0]
        snap = DaySnapshot()
        for kind, fname, attr in (
            ("open", "open_price.json", "open"),
            ("close", "close_price.json", "close"),
            ("vol", "vol_traded.json", "volume"),
        ):
            p = folder / fname
            if not p.exists():
                continue
            with p.open() as f:
                doc = json.load(f)
            cols = doc.get("columns", [])
            data = doc.get("data", [])
            if not data:
                continue
            row = data[0]
            for ci, company in enumerate(cols):
                if ci >= len(row):
                    continue
                v = _parse_archive_value(row[ci])
                existing = snap.rows.get(company)
                if existing is None:
                    existing = DailyRow(company, None, None, None)
                if v is not None:
                    existing = existing.with_field(attr, v)
                snap.rows[company] = existing
        if snap.rows:
            out[day].append(snap)
    return out


def merge_local_sources(
    folder_snaps: dict[str, list[DaySnapshot]],
    consol_data: dict[str, DaySnapshot],
) -> dict[str, tuple[DaySnapshot, list[DaySnapshot]]]:
    out: dict[str, tuple[DaySnapshot, list[DaySnapshot]]] = {}
    all_days = set(folder_snaps.keys()) | set(consol_data.keys())
    for day in all_days:
        snaps = list(folder_snaps.get(day, []))
        consol_snap = consol_data.get(day)

        primary: DaySnapshot | None = None
        duplicates: list[DaySnapshot] = []

        for snap in snaps:
            if primary is None:
                primary = snap
                continue
            for company, drow in snap.rows.items():
                prow = primary.rows.get(company)
                if prow is None:
                    primary.rows[company] = drow
                    continue
                for fld in FIELDS:
                    pv = getattr(prow, fld)
                    ev = getattr(drow, fld)
                    if ev is None or pv is None:
                        if ev is not None and pv is None:
                            primary.rows[company] = prow.with_field(fld, ev)
                        continue
                    if abs(pv - ev) > NUMERIC_TOLERANCE:
                        duplicates.append(DaySnapshot(rows={company: drow}))

        if consol_snap is not None:
            if primary is None:
                primary = consol_snap
            else:
                for company, crow in consol_snap.rows.items():
                    prow = primary.rows.get(company)
                    if prow is None:
                        primary.rows[company] = crow
                        continue
                    for fld in FIELDS:
                        cv = getattr(crow, fld)
                        pv = getattr(prow, fld)
                        if cv is None or pv is None:
                            continue
                        if abs(cv - pv) > NUMERIC_TOLERANCE:
                            duplicates.append(DaySnapshot(rows={company: crow}))

        if primary is None:
            primary = DaySnapshot()
        out[day] = (primary, duplicates)
    return out


def _rows_equal(a: DailyRow, b: DailyRow) -> bool:
    for av, bv in zip((a.open, a.close, a.volume), (b.open, b.close, b.volume)):
        if av is None and bv is None:
            continue
        if av is None or bv is None:
            return False
        if abs(av - bv) > NUMERIC_TOLERANCE:
            return False
    return True


def fmt(v: float | None) -> str:
    if v is None:
        return "missing"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return f"{v:.4f}".rstrip("0").rstrip(".")


def write_normalized_older(snapshots: dict[str, DaySnapshot], import_start: date, import_end: date) -> dict[str, Any]:
    in_window = {
        day: snap
        for day, snap in snapshots.items()
        if import_start <= date.fromisoformat(day) <= import_end
    }
    companies = sorted({c for s in in_window.values() for c in s.rows})
    days = sorted(in_window.keys())

    files: dict[str, Any] = {
        "opening_price.json": [],
        "closing_price.json": [],
        "vol_traded.json": [],
    }
    for field_name, fname, idx in (
        ("opening_price", "opening_price.json", 0),
        ("closing_price", "closing_price.json", 1),
        ("volume", "vol_traded.json", 2),
    ):
        out: dict[str, Any] = {"columns": companies, "index": days, "data": []}
        for day in days:
            snap = in_window[day]
            row: list[Any] = []
            for company in companies:
                drow = snap.rows.get(company)
                if drow is None:
                    row.append(None)
                else:
                    v = (drow.open, drow.close, drow.volume)[idx]
                    row.append(v)
            out["data"].append(row)
        target = NORMALIZED_DIR / fname
        with target.open("w") as f:
            json.dump(out, f, indent=2, default=str)
        files[fname] = out
    return {"companies": companies, "days": days, "in_window_count": len(days)}


def compare(
    older: dict[str, DaySnapshot],
    local: dict[str, tuple[DaySnapshot, list[DaySnapshot]]],
    comparison_start: date,
) -> list[Issue]:
    issues: list[Issue] = []
    all_days = set(older.keys()) | set(local.keys())
    in_window = sorted(d for d in all_days if date.fromisoformat(d) >= comparison_start)

    for day in in_window:
        older_snap = older.get(day)
        local_pair = local.get(day)
        local_snap = local_pair[0] if local_pair else None
        local_duplicates = local_pair[1] if local_pair else []

        older_dates_with_data = older_snap is not None and bool(older_snap.rows)
        local_dates_with_data = local_snap is not None and bool(local_snap.rows)

        if not older_dates_with_data and not local_dates_with_data:
            continue

        if not older_dates_with_data and local_dates_with_data:
            issues.append(Issue(day, "ALL", "ALL", "present", "missing", "missing_date_in_older_scraper",
                                note="Local archive has this date, but older scraper does not."))
        elif not local_dates_with_data and older_dates_with_data:
            issues.append(Issue(day, "ALL", "ALL", "missing", "present", "missing_date_in_local",
                                note="Older scraper has this date, but local archive does not."))
        else:
            assert older_snap is not None and local_snap is not None

            older_companies = set(older_snap.rows.keys())
            local_companies = set(local_snap.rows.keys())

            for company in sorted(local_companies - older_companies):
                for fld in FIELDS:
                    lv = getattr(local_snap.rows[company], fld)
                    issues.append(Issue(day, company, FIELD_DISPLAY[fld], fmt(lv), "missing", "missing_in_older_scraper"))
            for company in sorted(older_companies - local_companies):
                for fld in FIELDS:
                    ov = getattr(older_snap.rows[company], fld)
                    issues.append(Issue(day, company, FIELD_DISPLAY[fld], "missing", fmt(ov), "missing_in_local"))

            for company in sorted(older_companies & local_companies):
                for fld in FIELDS:
                    lv = getattr(local_snap.rows[company], fld)
                    ov = getattr(older_snap.rows[company], fld)
                    if lv is None and ov is None:
                        continue
                    if lv is None:
                        issues.append(Issue(day, company, FIELD_DISPLAY[fld], "missing", fmt(ov), "missing_in_local"))
                        continue
                    if ov is None:
                        issues.append(Issue(day, company, FIELD_DISPLAY[fld], fmt(lv), "missing", "missing_in_older_scraper"))
                        continue
                    if abs(lv - ov) > NUMERIC_TOLERANCE:
                        issues.append(Issue(day, company, FIELD_DISPLAY[fld], fmt(lv), fmt(ov), "numeric_difference"))

        if local_snap is not None and local_duplicates:
            seen: set[tuple[str, str]] = set()
            for dup_snap in local_duplicates:
                for company, drow in dup_snap.rows.items():
                    primary_row = local_snap.rows.get(company, drow)
                    for fld in FIELDS:
                        lv = getattr(primary_row, fld)
                        if lv is None:
                            continue
                        ov = getattr(drow, fld)
                        if ov is None or abs(lv - ov) > NUMERIC_TOLERANCE:
                            key = (company, fld)
                            if key in seen:
                                continue
                            seen.add(key)
                            issues.append(Issue(
                                day, company, FIELD_DISPLAY[fld], fmt(lv), fmt(ov),
                                "duplicate_in_local_archive",
                                note=f"local has another value: {fmt(ov)}",
                            ))

    return issues


def write_csv_report(issues: list[Issue], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["date", "company", "field", "local_value", "older_scraper_value", "issue_type", "note"])
        for issue in issues:
            writer.writerow([issue.date, issue.company, issue.field, issue.local_value,
                             issue.older_value, issue.issue_type, issue.note])


def _build_review_rows(issues: list[Issue]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for issue in issues:
        priority = "high" if issue.issue_type in ("missing_in_local", "duplicate_in_local_archive") else "normal"
        rows.append({
            "review_priority": priority,
            "date": issue.date,
            "company": issue.company,
            "field": issue.field,
            "this_repository_value": issue.local_value,
            "older_scraper_value": issue.older_value,
            "issue_type": issue.issue_type,
            "note": issue.note,
            "review_status": "pending",
            "review_decision": "",
            "review_notes": "",
        })
    return rows


def write_xlsx(
    issues: list[Issue],
    path: Path,
    sheet_prefix: str = "Review",
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = _build_review_rows(issues)

    by_type: dict[str, int] = defaultdict(int)
    by_type_field: dict[tuple[str, str], int] = defaultdict(int)
    for issue in issues:
        by_type[issue.issue_type] += 1
        by_type_field[(issue.issue_type, issue.field)] += 1

    numeric = [i for i in issues if i.issue_type == "numeric_difference"]
    missing = [i for i in issues if i.issue_type.startswith("missing_in_") and i.issue_type != "missing_date_in_"]
    missing_dates = [i for i in issues if i.issue_type.startswith("missing_date_in_")]

    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E5C8A")
    wrap = Alignment(wrap_text=True, vertical="top")

    def write_sheet(name: str, data_rows: list[dict[str, str]], title_cols: list[str]) -> None:
        ws = wb.create_sheet(title=name)
        for ci, col in enumerate(title_cols, start=1):
            cell = ws.cell(row=1, column=ci, value=col)
            cell.font = header_font
            cell.fill = header_fill
        for ri, row in enumerate(data_rows, start=2):
            for ci, col in enumerate(title_cols, start=1):
                c = ws.cell(row=ri, column=ci, value=row.get(col, ""))
                c.alignment = wrap
        widths = [16, 13, 42, 18, 22, 22, 28, 60, 18, 18, 28]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    review_cols = [
        "review_priority", "date", "company", "field",
        "this_repository_value", "older_scraper_value", "issue_type", "note",
        "review_status", "review_decision", "review_notes",
    ]
    write_sheet("Review List", rows, review_cols)
    write_sheet("Numeric Differences", _build_review_rows(numeric), review_cols)
    write_sheet("Missing Values", _build_review_rows(missing), review_cols)
    write_sheet("Missing Dates", _build_review_rows(missing_dates), review_cols)

    ws = wb.create_sheet(title="Summary")
    ws.cell(row=1, column=1, value="issue_type").font = header_font
    ws.cell(row=1, column=1).fill = header_fill
    ws.cell(row=1, column=2, value="count").font = header_font
    ws.cell(row=1, column=2).fill = header_fill
    for ri, (itype, count) in enumerate(sorted(by_type.items(), key=lambda kv: -kv[1]), start=2):
        ws.cell(row=ri, column=1, value=itype)
        ws.cell(row=ri, column=2, value=count)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 12

    ws = wb.create_sheet(title="Field Summary")
    for ci, h in enumerate(("issue_type", "field", "count"), start=1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = header_font
        c.fill = header_fill
    for ri, ((itype, fld), count) in enumerate(sorted(by_type_field.items()), start=2):
        ws.cell(row=ri, column=1, value=itype)
        ws.cell(row=ri, column=2, value=fld)
        ws.cell(row=ri, column=3, value=count)
    for i, w in enumerate((36, 18, 12), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    default = wb.active
    wb.remove(default)

    wb.save(path)


def main() -> None:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    NORMALIZED_DIR.mkdir(parents=True, exist_ok=True)

    print("Loading older scraper data from reference-data/older-zse/ ...")
    older = load_older_data()
    print(f"  days with data: {len(older)}")
    older_in_window = sum(1 for d in older if BRIEF_IMPORT_START <= date.fromisoformat(d) <= BRIEF_IMPORT_END)
    print(f"  days in import window 2020-10-08 to 2024-11-08: {older_in_window}")

    print("Loading local archive (folders + consolidated) ...")
    folder_data = load_local_from_folders()
    consol_data = load_local_from_consolidated()
    local = merge_local_sources(folder_data, consol_data)
    print(f"  local days: {len(local)}")

    print("Normalizing older scraper to archive-single-file shape ...")
    norm_meta = write_normalized_older(older, BRIEF_IMPORT_START, BRIEF_IMPORT_END)
    print(f"  companies: {len(norm_meta['companies'])}, days: {norm_meta['in_window_count']}")

    print("Comparing overlap from 2024-11-09 onward ...")
    issues = compare(older, local, BRIEF_LOCAL_START)
    print(f"  issues found: {len(issues)}")

    by_type: dict[str, int] = defaultdict(int)
    for issue in issues:
        by_type[issue.issue_type] += 1
    for itype, count in sorted(by_type.items(), key=lambda kv: -kv[1]):
        print(f"    {itype}: {count}")

    print("Writing reports ...")
    csv_path = REPORTS_DIR / "overlap_mismatch_report.csv"
    write_csv_report(issues, csv_path)
    print(f"  CSV: {csv_path.relative_to(REPO_ROOT)}")

    task4_path = REPORTS_DIR / "task4_review_workbook.xlsx"
    write_xlsx(issues, task4_path)
    print(f"  task4: {task4_path.relative_to(REPO_ROOT)}")

    task5_path = REPORTS_DIR / "task5_reviewable_mismatches_before_archive_changes.xlsx"
    write_xlsx(issues, task5_path)
    print(f"  task5: {task5_path.relative_to(REPO_ROOT)}")

    print("Done.")


if __name__ == "__main__":
    main()
